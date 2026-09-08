import io
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook


# ============================================================
# APP CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Pinestate Liquor Template Builder",
    page_icon="🍾",
    layout="wide",
)

APP_TITLE = "Pinestate Liquor Template Builder"

BASE_DIR = Path(__file__).parent
TEMPLATE_DIR = BASE_DIR / "templates"


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean(value):
    """
    Convert a value to a clean string.

    Blank/NaN values are returned as an empty string.
    """

    if value is None:
        return ""

    if pd.isna(value):
        return ""

    return str(value).strip()


def read_input_file(uploaded_file, header_row=0):
    """
    Read CSV or Excel input files.

    header_row:
        0 = headers are on Excel row 1
        1 = headers are on Excel row 2

    Promo Retail uses header_row=1 because:
        Row 2 = headers
        Row 3 = data
    """

    filename = uploaded_file.name.lower()

    if filename.endswith(".csv"):

        return pd.read_csv(
            uploaded_file,
            dtype=str,
            keep_default_na=False,
            header=header_row,
        )

    elif filename.endswith((".xlsx", ".xls")):

        return pd.read_excel(
            uploaded_file,
            dtype=str,
            keep_default_na=False,
            header=header_row,
        )

    else:

        raise ValueError(
            f"Unsupported file type: {uploaded_file.name}"
        )


def source_column(df, excel_column):
    """
    Retrieve a DataFrame column using Excel-style letters.

    A = first column
    B = second column
    H = eighth column
    O = fifteenth column
    """

    index = ord(excel_column.upper()) - ord("A")

    if index < 0 or index >= len(df.columns):

        raise ValueError(
            f"Input file does not contain column {excel_column}."
        )

    return df.iloc[:, index].map(clean)


def load_template(filename):
    """
    Load an Excel template from the templates folder.
    """

    path = TEMPLATE_DIR / filename

    if not path.exists():

        raise FileNotFoundError(
            f"Template '{filename}' was not found in the templates folder."
        )

    return load_workbook(path)


def write_to_template(
    template_filename,
    rows,
    start_row,
):
    """
    Write generated records into an Excel template.

    start_row determines where generated data begins.

    Promo Retail:
        Data starts row 2

    Standard Cost:
        Data starts row 12

    Promo Cost:
        Data starts row 12
    """

    wb = load_template(template_filename)

    # Use the first worksheet in the supplied template.
    ws = wb[wb.sheetnames[0]]

    # --------------------------------------------------------
    # Clear existing data below the header.
    # --------------------------------------------------------

    if ws.max_row >= start_row:

        for row in ws.iter_rows(
            min_row=start_row,
            max_row=ws.max_row,
            min_col=1,
            max_col=ws.max_column,
        ):

            for cell in row:
                cell.value = None

    # --------------------------------------------------------
    # Write generated records.
    # --------------------------------------------------------

    for row_number, row_data in enumerate(
        rows,
        start=start_row,
    ):

        for column_number, value in enumerate(
            row_data,
            start=1,
        ):

            ws.cell(
                row=row_number,
                column=column_number,
            ).value = value

    # --------------------------------------------------------
    # Save workbook into memory.
    # --------------------------------------------------------

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()


def remove_duplicates(rows):
    """
    Remove exact duplicate output records.

    Original order is preserved.
    """

    seen = set()

    unique_rows = []

    for row in rows:

        key = tuple(row)

        if key not in seen:

            seen.add(key)

            unique_rows.append(row)

    return unique_rows


# ============================================================
# PRODUCTS LOOKUP
# ============================================================

def build_products_lookup(products):
    """
    Products File lookup:

        Products Column D = lookup key
        Products Column H = value returned to Output Column B
    """

    lookup = {}

    products_key = source_column(
        products,
        "D",
    )

    products_value = source_column(
        products,
        "H",
    )

    for key, value in zip(
        products_key,
        products_value,
    ):

        key = clean(key)

        if key and key not in lookup:

            lookup[key] = clean(value)

    return lookup


# ============================================================
# EG PROMO RETAIL
# ============================================================

def build_eg_promo_retail(
    promo_retail,
):
    """
    EG Promo Retail mapping.

    Source:
        Promo Retail File

    Output:

        C = Source B
        F = Source C
        H = Source E
        K = Source F
        L = Source I
        N = 08784
        O = Pine State Liquor EG
        P = 0

    All other output columns remain blank.
    """

    rows = []

    for _, record in promo_retail.iterrows():

        # Output has columns A:P = 16 columns.
        output = [""] * 16

        # ----------------------------------------------------
        # C <- Promo Retail Column B
        # ----------------------------------------------------

        output[2] = clean(
            record.iloc[1]
        )

        # ----------------------------------------------------
        # F <- Promo Retail Column C
        # ----------------------------------------------------

        output[5] = clean(
            record.iloc[2]
        )

        # ----------------------------------------------------
        # H <- Promo Retail Column E
        # ----------------------------------------------------

        output[7] = clean(
            record.iloc[4]
        )

        # ----------------------------------------------------
        # K <- Promo Retail Column F
        # ----------------------------------------------------

        output[10] = clean(
            record.iloc[5]
        )

        # ----------------------------------------------------
        # L <- Promo Retail Column I
        # ----------------------------------------------------

        output[11] = clean(
            record.iloc[8]
        )

        # ----------------------------------------------------
        # N = Default Vendor ID
        # ----------------------------------------------------

        output[13] = "08784"

        # ----------------------------------------------------
        # O = Default Vendor Description
        # ----------------------------------------------------

        output[14] = "Pine State Liquor EG"

        # ----------------------------------------------------
        # P = Default Cost Zone
        # ----------------------------------------------------

        output[15] = "0"

        rows.append(output)

    # --------------------------------------------------------
    # Remove duplicate records.
    # --------------------------------------------------------

    return remove_duplicates(rows)


# ============================================================
# EG STANDARD COST
# ============================================================

def build_eg_standard_cost(
    raw_cost,
    products,
):
    """
    EG Standard Cost.

    ALL records from the Raw Vendor Store Cost File
    are included.

    There is NO filter on Raw Column L.

    Output mapping:

        A = VC
        B = Products Column H
        C = Raw Column C
        D = Raw Column O
        E = blank
        F = Raw Column K
        G = Raw Column M
        H = blank
        I = Raw Column B
        J = blank
        K = blank
        L = blank

    Lookup:

        Output I
            ↓
        Products Column D
            ↓
        Products Column H
            ↓
        Output B

    If there is no match:
        Output B = blank
    """

    products_lookup = build_products_lookup(
        products
    )

    rows = []

    # --------------------------------------------------------
    # IMPORTANT:
    #
    # ALL raw cost records are processed.
    #
    # There is intentionally NO:
    #
    #     Column L = 0
    #
    # filter here.
    # --------------------------------------------------------

    for _, record in raw_cost.iterrows():

        # Output A:L = 12 columns.
        output = [""] * 12

        # ----------------------------------------------------
        # A = VC
        # ----------------------------------------------------

        output[0] = "VC"

        # ----------------------------------------------------
        # C <- Raw Column C
        # ----------------------------------------------------

        output[2] = clean(
            record.iloc[2]
        )

        # ----------------------------------------------------
        # D <- Raw Column O
        # ----------------------------------------------------

        output[3] = clean(
            record.iloc[14]
        )

        # ----------------------------------------------------
        # F <- Raw Column K
        # ----------------------------------------------------

        output[5] = clean(
            record.iloc[10]
        )

        # ----------------------------------------------------
        # G <- Raw Column M
        # ----------------------------------------------------

        output[6] = clean(
            record.iloc[12]
        )

        # ----------------------------------------------------
        # I <- Raw Column B
        # ----------------------------------------------------

        output[8] = clean(
            record.iloc[1]
        )

        # ----------------------------------------------------
        # B LOOKUP
        #
        # Output I
        #      ↓
        # Products D
        #      ↓
        # Products H
        #      ↓
        # Output B
        #
        # No match = blank.
        # ----------------------------------------------------

        lookup_key = output[8]

        if lookup_key:

            output[1] = products_lookup.get(
                lookup_key,
                "",
            )

        rows.append(output)

    # --------------------------------------------------------
    # Remove duplicate records.
    # --------------------------------------------------------

    return remove_duplicates(rows)


# ============================================================
# EG PROMO COST
# ============================================================

def build_eg_promo_cost(
    raw_cost,
    products,
):
    """
    EG Promo Cost.

    ONLY records where:

        Raw Vendor Store Cost Column L = 1

    are included.

    Output mapping:

        A = VC
        B = Products Column H
        C = Raw Column C
        D = Raw Column O
        E = Raw Column L
        F = Raw Column K
        G = Raw Column M
        H = Raw Column N
        I = Raw Column B
        J = blank
        K = blank
        L = blank

    Lookup:

        Output I
            ↓
        Products Column D
            ↓
        Products Column H
            ↓
        Output B

    If there is no match:
        Output B = blank
    """

    products_lookup = build_products_lookup(
        products
    )

    rows = []

    for _, record in raw_cost.iterrows():

        # ----------------------------------------------------
        # Raw Column L
        # ----------------------------------------------------

        promo_flag = clean(
            record.iloc[11]
        )

        # ----------------------------------------------------
        # Promo Cost ONLY uses Column L = 1.
        # ----------------------------------------------------

        if promo_flag != "1":

            continue

        # Output A:L = 12 columns.
        output = [""] * 12

        # ----------------------------------------------------
        # A = VC
        # ----------------------------------------------------

        output[0] = "VC"

        # ----------------------------------------------------
        # C <- Raw Column C
        # ----------------------------------------------------

        output[2] = clean(
            record.iloc[2]
        )

        # ----------------------------------------------------
        # D <- Raw Column O
        # ----------------------------------------------------

        output[3] = clean(
            record.iloc[14]
        )

        # ----------------------------------------------------
        # E <- Raw Column L
        # ----------------------------------------------------

        output[4] = clean(
            record.iloc[11]
        )

        # ----------------------------------------------------
        # F <- Raw Column K
        # ----------------------------------------------------

        output[5] = clean(
            record.iloc[10]
        )

        # ----------------------------------------------------
        # G <- Raw Column M
        # ----------------------------------------------------

        output[6] = clean(
            record.iloc[12]
        )

        # ----------------------------------------------------
        # H <- Raw Column N
        # ----------------------------------------------------

        output[7] = clean(
            record.iloc[13]
        )

        # ----------------------------------------------------
        # I <- Raw Column B
        # ----------------------------------------------------

        output[8] = clean(
            record.iloc[1]
        )

        # ----------------------------------------------------
        # B LOOKUP
        #
        # Output I
        #      ↓
        # Products D
        #      ↓
        # Products H
        #      ↓
        # Output B
        #
        # No match = blank.
        # ----------------------------------------------------

        lookup_key = output[8]

        if lookup_key:

            output[1] = products_lookup.get(
                lookup_key,
                "",
            )

        rows.append(output)

    # --------------------------------------------------------
    # Remove duplicate records.
    # --------------------------------------------------------

    return remove_duplicates(rows)


# ============================================================
# STREAMLIT USER INTERFACE
# ============================================================

st.title(APP_TITLE)

st.markdown(
    """
    Upload the three required source files to generate
    the Pine State Liquor EG templates.
    """
)


# ============================================================
# FILE UPLOADERS
# ============================================================

st.subheader("Required Files")

col1, col2, col3 = st.columns(3)


with col1:

    products_file = st.file_uploader(
        "1. Products File",
        type=[
            "csv",
            "xlsx",
            "xls",
        ],
        key="products_file",
    )


with col2:

    promo_retail_file = st.file_uploader(
        "2. Promo Retail File",
        type=[
            "csv",
            "xlsx",
            "xls",
        ],
        key="promo_retail_file",
    )


with col3:

    raw_cost_file = st.file_uploader(
        "3. Raw Vendor Store Cost File",
        type=[
            "csv",
            "xlsx",
            "xls",
        ],
        key="raw_cost_file",
    )


st.divider()


# ============================================================
# PROCESS BUTTON
# ============================================================

process = st.button(
    "🚀 Process Files",
    type="primary",
    use_container_width=True,
)


if process:

    # ========================================================
    # VALIDATE UPLOADS
    # ========================================================

    if not products_file:

        st.error(
            "Please upload the Products File."
        )

        st.stop()


    if not promo_retail_file:

        st.error(
            "Please upload the Promo Retail File."
        )

        st.stop()


    if not raw_cost_file:

        st.error(
            "Please upload the Raw Vendor Store Cost File."
        )

        st.stop()


    try:

        with st.spinner(
            "Processing files..."
        ):

            # =================================================
            # READ INPUT FILES
            # =================================================

            # Products:
            # Row 1 = headers
            # Row 2 onward = data

            products = read_input_file(
                products_file,
                header_row=0,
            )


            # Promo Retail:
            # Row 2 = headers
            # Row 3 onward = data

            promo_retail = read_input_file(
                promo_retail_file,
                header_row=1,
            )


            # Raw Vendor Store Cost:
            # Row 1 = headers
            # Row 2 onward = data

            raw_cost = read_input_file(
                raw_cost_file,
                header_row=0,
            )


            # =================================================
            # VALIDATE SOURCE COLUMN COUNTS
            # =================================================

            if len(products.columns) < 8:

                raise ValueError(
                    "Products File must contain at least 8 columns."
                )


            if len(promo_retail.columns) < 9:

                raise ValueError(
                    "Promo Retail File must contain at least 9 columns."
                )


            if len(raw_cost.columns) < 15:

                raise ValueError(
                    "Raw Vendor Store Cost File must contain at least 15 columns."
                )


            # =================================================
            # BUILD OUTPUT DATA
            # =================================================

            promo_retail_rows = build_eg_promo_retail(
                promo_retail
            )


            standard_cost_rows = build_eg_standard_cost(
                raw_cost,
                products,
            )


            promo_cost_rows = build_eg_promo_cost(
                raw_cost,
                products,
            )


            # =================================================
            # WRITE OUTPUT TEMPLATES
            # =================================================

            # -------------------------------------------------
            # EG PROMO RETAIL
            #
            # Output data starts at row 2.
            # -------------------------------------------------

            promo_retail_output = write_to_template(
                "EG_PromoRetail.xlsx",
                promo_retail_rows,
                start_row=2,
            )


            # -------------------------------------------------
            # EG STANDARD COST
            #
            # Header = row 11
            # Data = row 12
            # -------------------------------------------------

            standard_cost_output = write_to_template(
                "EG_StandardCost.xlsx",
                standard_cost_rows,
                start_row=12,
            )


            # -------------------------------------------------
            # EG PROMO COST
            #
            # Header = row 11
            # Data = row 12
            # -------------------------------------------------

            promo_cost_output = write_to_template(
                "EG_PromoCost.xlsx",
                promo_cost_rows,
                start_row=12,
            )


            # =================================================
            # SAVE OUTPUTS IN SESSION STATE
            # =================================================

            st.session_state[
                "promo_retail_output"
            ] = promo_retail_output


            st.session_state[
                "standard_cost_output"
            ] = standard_cost_output


            st.session_state[
                "promo_cost_output"
            ] = promo_cost_output


            st.session_state[
                "promo_retail_count"
            ] = len(
                promo_retail_rows
            )


            st.session_state[
                "standard_cost_count"
            ] = len(
                standard_cost_rows
            )


            st.session_state[
                "promo_cost_count"
            ] = len(
                promo_cost_rows
            )


            st.session_state[
                "processed"
            ] = True


        # =====================================================
        # SUCCESS
        # =====================================================

        st.success(
            "Files processed successfully!"
        )


    except Exception as error:

        st.error(
            f"Processing failed: {error}"
        )

        st.stop()


# ============================================================
# DOWNLOAD SECTION
# ============================================================

if st.session_state.get(
    "processed",
    False,
):

    st.subheader(
        "Generated Files"
    )


    # ========================================================
    # EG PROMO RETAIL
    # ========================================================

    col1, col2 = st.columns(
        [4, 1]
    )


    with col1:

        st.write(
            f"""
            **EG_Promo Retail File.xlsx**

            {st.session_state['promo_retail_count']:,} records
            """
        )


    with col2:

        st.download_button(
            label="Download",
            data=st.session_state[
                "promo_retail_output"
            ],
            file_name=(
                "EG_Promo Retail File.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            key="download_promo_retail",
            use_container_width=True,
        )


    # ========================================================
    # EG STANDARD COST
    # ========================================================

    col1, col2 = st.columns(
        [4, 1]
    )


    with col1:

        st.write(
            f"""
            **EG_Standard Cost File.xlsx**

            {st.session_state['standard_cost_count']:,} records
            """
        )


    with col2:

        st.download_button(
            label="Download",
            data=st.session_state[
                "standard_cost_output"
            ],
            file_name=(
                "EG_Standard Cost File.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            key="download_standard_cost",
            use_container_width=True,
        )


    # ========================================================
    # EG PROMO COST
    # ========================================================

    col1, col2 = st.columns(
        [4, 1]
    )


    with col1:

        st.write(
            f"""
            **EG_Promo Cost File.xlsx**

            {st.session_state['promo_cost_count']:,} records
            """
        )


    with col2:

        st.download_button(
            label="Download",
            data=st.session_state[
                "promo_cost_output"
            ],
            file_name=(
                "EG_Promo Cost File.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            key="download_promo_cost",
            use_container_width=True,
        )


    st.divider()


    st.caption(
        "Duplicate records are automatically removed "
        "before the output files are generated."
    )
